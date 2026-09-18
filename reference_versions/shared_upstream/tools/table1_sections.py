from __future__ import annotations

import re
import unicodedata
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


TITLE_RE = re.compile(r"(『.*?』)", re.DOTALL)


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour == value.minute == value.second == value.microsecond == 0:
            text = value.strftime("%Y-%m-%d")
        elif value.second == value.microsecond == 0:
            text = value.strftime("%Y-%m-%d %H:%M")
        else:
            text = value.strftime("%Y-%m-%d %H:%M:%S")
    elif isinstance(value, date):
        text = value.strftime("%Y-%m-%d")
    elif isinstance(value, time):
        text = value.strftime("%H:%M:%S" if value.second else "%H:%M")
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = unicodedata.normalize("NFKC", text.replace("\r\n", "\n").replace("\r", "\n"))
    return "\n".join(line.rstrip() for line in text.splitlines()).strip()


@dataclass(frozen=True)
class GridCell:
    row: int
    col: int
    x: float
    text: str


@dataclass(frozen=True)
class GridRow:
    number: int
    cells: tuple[GridCell, ...]

    @property
    def text(self) -> str:
        return "\n".join(cell.text for cell in self.cells if cell.text)


@dataclass
class TitledSection:
    title: str
    fragments: list[tuple[int, str]] = field(default_factory=list)

    @property
    def row_numbers(self) -> list[int]:
        return sorted({row for row, text in self.fragments if text.strip()})

    @property
    def text(self) -> str:
        return "\n".join(text for _, text in self.fragments if text.strip()).strip()

    @property
    def start_row(self) -> int:
        rows = self.row_numbers
        return rows[0] if rows else 0

    @property
    def end_row(self) -> int:
        rows = self.row_numbers
        return rows[-1] if rows else 0

    def add(self, row: int, text: str) -> None:
        cleaned = text.strip()
        if cleaned:
            self.fragments.append((row, cleaned))


@dataclass(frozen=True)
class Table1Grid:
    source_path: Path
    sheet_name: str
    rows: tuple[GridRow, ...]
    max_col: int

    def row(self, number: int) -> GridRow:
        if number <= 0 or number >= len(self.rows):
            return GridRow(number=number, cells=())
        return self.rows[number]

    def iter_rows(self, start: int = 1, end: int | None = None) -> Iterable[GridRow]:
        upper = min(end or (len(self.rows) - 1), len(self.rows) - 1)
        for number in range(max(1, start), upper + 1):
            yield self.rows[number]


def _merge_centers(sheet: Any) -> dict[tuple[int, int], float]:
    centers: dict[tuple[int, int], float] = {}
    for merged in sheet.merged_cells.ranges:
        centers[(merged.min_row, merged.min_col)] = (merged.min_col + merged.max_col - 2) / 2
    return centers


def load_table1_grid(excel_path: Path, sheet_name: str) -> Table1Grid:
    excel_path = excel_path.expanduser().resolve()
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="DrawingML support is incomplete.*")
        workbook = load_workbook(excel_path, read_only=False, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            available = ", ".join(workbook.sheetnames)
            raise ValueError(f"Sheet {sheet_name!r} not found. Available sheets: {available}")
        sheet = workbook[sheet_name]
        centers = _merge_centers(sheet)
        rows: list[GridRow] = [GridRow(number=0, cells=())]
        for row_number in range(1, sheet.max_row + 1):
            cells: list[GridCell] = []
            for col_number in range(1, sheet.max_column + 1):
                text = normalize_text(sheet.cell(row_number, col_number).value)
                if not text:
                    continue
                cells.append(
                    GridCell(
                        row=row_number,
                        col=col_number - 1,
                        x=centers.get((row_number, col_number), float(col_number - 1)),
                        text=text,
                    )
                )
            rows.append(GridRow(number=row_number, cells=tuple(cells)))
        return Table1Grid(
            source_path=excel_path,
            sheet_name=sheet_name,
            rows=tuple(rows),
            max_col=sheet.max_column,
        )
    finally:
        workbook.close()


def clean_section_title(raw_title: str) -> str:
    title = normalize_text(raw_title).strip("『』 ")
    title = re.sub(r"項目/日期", "", title)
    title = re.sub(r"\s+", " ", title).strip()
    return title


def split_titled_sections(grid: Table1Grid, *, end_row: int | None = None) -> list[TitledSection]:
    sections: list[TitledSection] = []
    current: TitledSection | None = None
    for row in grid.iter_rows(end=end_row):
        parts = TITLE_RE.split(row.text)
        for part in parts:
            if not part or not part.strip():
                continue
            if part.startswith("『") and part.endswith("』"):
                if current is not None:
                    sections.append(current)
                current = TitledSection(title=clean_section_title(part))
                current.add(row.number, part)
            elif current is not None:
                current.add(row.number, part)
    if current is not None:
        sections.append(current)
    return sections


def find_row(grid: Table1Grid, needle: str, *, start: int = 1) -> int | None:
    normalized = normalize_text(needle).casefold()
    for row in grid.iter_rows(start=start):
        if normalized in row.text.casefold():
            return row.number
    return None
