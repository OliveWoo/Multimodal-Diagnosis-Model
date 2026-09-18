from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE


@dataclass
class ExtractedImage:
    image_index: int
    raw_path: Path
    width: int
    height: int
    ext: str
    shape_name: str
    shape_order: int
    display_geometry: dict[str, int]


@dataclass
class SlideExtraction:
    slide_index: int
    text_blocks: list[str]
    text_records: list[dict[str, Any]]
    images: list[ExtractedImage]
    debug_shapes: list[dict[str, Any]]


def _shape_geometry(shape, *, shape_order: int) -> dict[str, int]:
    left = int(getattr(shape, "left", 0) or 0)
    top = int(getattr(shape, "top", 0) or 0)
    width = int(getattr(shape, "width", 0) or 0)
    height = int(getattr(shape, "height", 0) or 0)
    return {
        "shape_order": shape_order,
        "left": left,
        "top": top,
        "width": width,
        "height": height,
        "right": left + width,
        "bottom": top + height,
        "center_x": left + (width // 2),
        "center_y": top + (height // 2),
    }


def _collect_text(
    shape,
    text_blocks: list[str],
    text_records: list[dict[str, Any]],
    debug_shapes: list[dict[str, Any]],
    *,
    shape_order: int,
) -> None:
    if getattr(shape, "has_text_frame", False):
        text = shape.text.strip()
        if text:
            text_blocks.append(text)
            geometry = _shape_geometry(shape, shape_order=shape_order)
            text_record = {
                "text": text,
                "shape_name": shape.name,
                "geometry": geometry,
            }
            text_records.append(text_record)
            debug_shapes.append(
                {
                    "type": "text",
                    "name": shape.name,
                    "text": text,
                    "geometry": geometry,
                }
            )


def _extract_picture(shape, image_index: int, images_dir: Path, *, shape_order: int) -> ExtractedImage:
    image = shape.image
    ext = image.ext.lower()
    raw_path = images_dir / f"raw_{image_index:02d}.{ext}"
    raw_path.write_bytes(image.blob)
    return ExtractedImage(
        image_index=image_index,
        raw_path=raw_path,
        width=image.size[0],
        height=image.size[1],
        ext=ext,
        shape_name=shape.name,
        shape_order=shape_order,
        display_geometry=_shape_geometry(shape, shape_order=shape_order),
    )


def _walk_shapes(
    shapes,
    images_dir: Path,
    text_blocks: list[str],
    text_records: list[dict[str, Any]],
    images: list[ExtractedImage],
    debug_shapes: list[dict[str, Any]],
    shape_order_ref: list[int],
) -> None:
    for shape in shapes:
        shape_order_ref[0] += 1
        shape_order = shape_order_ref[0]
        _collect_text(
            shape,
            text_blocks,
            text_records,
            debug_shapes,
            shape_order=shape_order,
        )
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image_index = len(images) + 1
            image = _extract_picture(shape, image_index, images_dir, shape_order=shape_order)
            images.append(image)
            debug_shapes.append(
                {
                    "type": "picture",
                    "name": shape.name,
                    "image_index": image.image_index,
                    "raw_path": str(image.raw_path),
                    "width": image.width,
                    "height": image.height,
                    "geometry": image.display_geometry,
                }
            )
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            debug_shapes.append(
                {
                    "type": "group",
                    "name": shape.name,
                    "geometry": _shape_geometry(shape, shape_order=shape_order),
                }
            )
            _walk_shapes(
                shape.shapes,
                images_dir,
                text_blocks,
                text_records,
                images,
                debug_shapes,
                shape_order_ref,
            )


def extract_slide_payload(slide, *, slide_index: int, images_dir: Path) -> SlideExtraction:
    text_blocks: list[str] = []
    text_records: list[dict[str, Any]] = []
    images: list[ExtractedImage] = []
    debug_shapes: list[dict[str, Any]] = []
    shape_order_ref = [0]
    _walk_shapes(
        slide.shapes,
        images_dir,
        text_blocks,
        text_records,
        images,
        debug_shapes,
        shape_order_ref,
    )
    return SlideExtraction(
        slide_index=slide_index,
        text_blocks=text_blocks,
        text_records=text_records,
        images=images,
        debug_shapes=debug_shapes,
    )
