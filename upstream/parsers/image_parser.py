from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Any, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.formula import ArrayFormula

from utils import sanitize_filename

SHEET_NAME = "image"


def _ensure_dir(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path


def _normalize_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _sheet_to_rows(
    sheet_values: Worksheet,
    sheet_formulas: Worksheet,
) -> Tuple[list[list[str]], dict[str, str]]:
    rows: list[list[str]] = []
    formulas: dict[str, str] = {}

    for row_formula, row_value in zip(
        sheet_formulas.iter_rows(),
        sheet_values.iter_rows(values_only=True),
    ):
        value_list = list(row_value)
        row_data: list[str] = []
        for index, cell_formula in enumerate(row_formula):
            cell_value = value_list[index] if index < len(value_list) else None
            if cell_formula.data_type == "f" and cell_formula.value is not None:
                raw_formula = cell_formula.value
                if isinstance(raw_formula, ArrayFormula):
                    formula_text = raw_formula.text
                else:
                    formula_text = str(raw_formula)
                row_data.append(formula_text)
                formulas[cell_formula.coordinate] = formula_text
            else:
                row_data.append(_normalize_value(cell_value))
        rows.append(row_data)

    return rows, formulas


def parse(excel_path: Path | str, output_dir: Path | str) -> Path:
    excel_path = Path(excel_path)
    output_dir = _ensure_dir(Path(output_dir))

    workbook_values = load_workbook(excel_path, data_only=True, read_only=True)
    workbook_formulas = load_workbook(excel_path, data_only=False, read_only=True)

    if SHEET_NAME not in workbook_values.sheetnames:
        workbook_values.close()
        workbook_formulas.close()
        raise ValueError(f"找不到工作表：{SHEET_NAME}")

    sheet_values = workbook_values[SHEET_NAME]
    sheet_formulas = workbook_formulas[SHEET_NAME]
    raw_data, formulas = _sheet_to_rows(sheet_values, sheet_formulas)

    workbook_values.close()
    workbook_formulas.close()

    payload = {
        "source_file": excel_path.name,
        "sheet": SHEET_NAME,
        "raw_data": raw_data,
    }
    if formulas:
        payload["formulas"] = formulas

    sanitized_workbook = sanitize_filename(excel_path.stem)
    sanitized_sheet = sanitize_filename(SHEET_NAME)
    raw_filename = f"{sanitized_workbook}_{sanitized_sheet}_Raw.json"
    output_path = output_dir / raw_filename
    with output_path.open("w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)

    return output_path


__all__ = ["SHEET_NAME", "parse"]