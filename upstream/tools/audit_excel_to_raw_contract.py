"""Audit available Excel workbooks against historical per-sheet Raw JSON."""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Sequence

from openpyxl import load_workbook

from tools.audit_source_to_normalized_contract import text_integrity_summary


AUDIT_VERSION = "excel_to_raw_contract_v1"


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    formula_text = getattr(value, "text", None)
    if isinstance(formula_text, str):
        return formula_text
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def canonical_text(value: Any) -> str:
    text = unicodedata.normalize("NFKC", cell_text(value))
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def loose_text(value: Any) -> str:
    return re.sub(r"\s+", "", canonical_text(value)).casefold()


def semantic_excel_candidates(data_value: Any, formula_value: Any) -> set[str]:
    candidates = {canonical_text(data_value), canonical_text(formula_value)}
    for value in (data_value, formula_value):
        if isinstance(value, datetime):
            if value.time().isoformat() == "00:00:00":
                candidates.add(value.date().isoformat())
        elif hasattr(value, "hour") and hasattr(value, "minute"):
            if getattr(value, "second", None) == 0:
                candidates.add(f"{value.hour:02d}:{value.minute:02d}")
    return {candidate for candidate in candidates if candidate}


def workbook_index(roots: Sequence[Path]) -> dict[str, list[Path]]:
    result: dict[str, list[Path]] = {}
    for root in roots:
        paths: list[Path] = []
        if root.is_file() and root.suffix.lower() in {".xlsx", ".xlsm"}:
            paths = [root]
        elif root.is_dir():
            paths = [*root.rglob("*.xlsx"), *root.rglob("*.xlsm")]
        for path in paths:
            result.setdefault(path.name.casefold(), []).append(path)
    return result


def compare_sheet(raw_path: Path, workbook_path: Path) -> dict[str, Any]:
    raw = read_json(raw_path)
    sheet_name = str(raw.get("sheet") or "")
    rows = raw.get("raw_data") or []
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    formula_workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    sheet_lookup = {name.casefold(): name for name in workbook.sheetnames}
    actual_sheet_name = sheet_lookup.get(sheet_name.casefold())
    if actual_sheet_name is None:
        workbook.close()
        formula_workbook.close()
        return {
            "raw_path": str(raw_path),
            "workbook_path": str(workbook_path),
            "sheet": sheet_name,
            "status": "worksheet_missing",
            "available_sheets": workbook.sheetnames,
        }

    worksheet = workbook[actual_sheet_name]
    formula_worksheet = formula_workbook[actual_sheet_name]
    workbook_dimensions = [worksheet.max_row, worksheet.max_column]
    raw_row_count = len(rows)
    raw_column_count = max((len(row) for row in rows if isinstance(row, list)), default=0)
    exact_matches = 0
    loose_matches = 0
    semantic_matches = 0
    compared_cells = 0
    raw_nonempty = 0
    workbook_nonempty = 0
    raw_missing = 0
    value_mismatches = 0
    mismatch_examples: list[dict[str, Any]] = []
    workbook_values: list[str] = []

    for row_index in range(1, raw_row_count + 1):
        raw_row = rows[row_index - 1] if isinstance(rows[row_index - 1], list) else []
        for column_index in range(1, raw_column_count + 1):
            raw_value = raw_row[column_index - 1] if column_index <= len(raw_row) else ""
            workbook_value = worksheet.cell(row_index, column_index).value
            formula_value = formula_worksheet.cell(row_index, column_index).value
            raw_text = canonical_text(raw_value)
            workbook_text = canonical_text(workbook_value)
            workbook_values.append(workbook_text)
            if raw_text:
                raw_nonempty += 1
            if workbook_text:
                workbook_nonempty += 1
            if not raw_text and not workbook_text:
                continue
            compared_cells += 1
            if raw_text == workbook_text:
                exact_matches += 1
                loose_matches += 1
                semantic_matches += 1
                continue
            if loose_text(raw_text) == loose_text(workbook_text):
                loose_matches += 1
                semantic_matches += 1
                continue
            if raw_text in semantic_excel_candidates(workbook_value, formula_value):
                semantic_matches += 1
                continue
            if workbook_text and not raw_text:
                raw_missing += 1
            else:
                value_mismatches += 1
            if len(mismatch_examples) < 12:
                mismatch_examples.append(
                    {
                        "cell": worksheet.cell(row_index, column_index).coordinate,
                        "excel": workbook_text[:300],
                        "raw": raw_text[:300],
                    }
                )

    outside_raw_bounds_nonempty = 0
    outside_raw_bounds_examples: list[dict[str, Any]] = []
    for row_index, row in enumerate(formula_worksheet.iter_rows(), start=1):
        for column_index, cell in enumerate(row, start=1):
            if row_index <= raw_row_count and column_index <= raw_column_count:
                continue
            value = canonical_text(cell.value)
            if not value:
                continue
            outside_raw_bounds_nonempty += 1
            if len(outside_raw_bounds_examples) < 12:
                outside_raw_bounds_examples.append(
                    {
                        "cell": f"R{row_index}C{column_index}",
                        "excel": value[:300],
                    }
                )

    workbook_integrity = text_integrity_summary(workbook_values)
    raw_integrity = text_integrity_summary(raw)
    workbook.close()
    formula_workbook.close()
    return {
        "raw_path": str(raw_path),
        "workbook_path": str(workbook_path),
        "sheet": sheet_name,
        "status": "compared",
        "raw_dimensions": [raw_row_count, raw_column_count],
        "workbook_dimensions": workbook_dimensions,
        "compared_nonempty_union_cells": compared_cells,
        "exact_match_count": exact_matches,
        "loose_match_count": loose_matches,
        "semantic_match_count": semantic_matches,
        "raw_nonempty_count": raw_nonempty,
        "workbook_nonempty_count_within_raw_bounds": workbook_nonempty,
        "raw_missing_count": raw_missing,
        "value_mismatch_count": value_mismatches,
        "outside_raw_bounds_nonempty_count": outside_raw_bounds_nonempty,
        "outside_raw_bounds_examples": outside_raw_bounds_examples,
        "mismatch_examples": mismatch_examples,
        "excel_text_integrity": workbook_integrity,
        "raw_text_integrity": raw_integrity,
    }


def audit(raw_root: Path, workbook_roots: Sequence[Path]) -> dict[str, Any]:
    index = workbook_index(workbook_roots)
    results: list[dict[str, Any]] = []
    unresolved: set[str] = set()
    for raw_path in sorted(raw_root.glob("NGS_patient_*_json_Raw/*_Raw.json")):
        raw = read_json(raw_path)
        source_file = str(raw.get("source_file") or "")
        matches = index.get(source_file.casefold(), [])
        if not matches:
            unresolved.add(source_file)
            continue
        results.append(compare_sheet(raw_path, matches[0]))

    status_counts = Counter(item["status"] for item in results)
    compared = [item for item in results if item["status"] == "compared"]
    return {
        "audit_version": AUDIT_VERSION,
        "answer_blind": True,
        "tier_effect": "none_audit_only",
        "raw_root": str(raw_root),
        "workbook_roots": [str(path) for path in workbook_roots],
        "resolved_sheet_count": len(results),
        "resolved_workbook_count": len(
            {item["workbook_path"] for item in results if item.get("workbook_path")}
        ),
        "status_counts": dict(status_counts),
        "exact_match_count": sum(item.get("exact_match_count", 0) for item in compared),
        "loose_match_count": sum(item.get("loose_match_count", 0) for item in compared),
        "semantic_match_count": sum(
            item.get("semantic_match_count", 0) for item in compared
        ),
        "compared_nonempty_union_cells": sum(
            item.get("compared_nonempty_union_cells", 0) for item in compared
        ),
        "raw_missing_count": sum(item.get("raw_missing_count", 0) for item in compared),
        "value_mismatch_count": sum(
            item.get("value_mismatch_count", 0) for item in compared
        ),
        "outside_raw_bounds_nonempty_count": sum(
            item.get("outside_raw_bounds_nonempty_count", 0) for item in compared
        ),
        "excel_suspicious_sheet_count": sum(
            bool((item.get("excel_text_integrity") or {}).get("has_suspicious_text"))
            for item in compared
        ),
        "raw_suspicious_sheet_count": sum(
            bool((item.get("raw_text_integrity") or {}).get("has_suspicious_text"))
            for item in compared
        ),
        "unresolved_source_workbooks": sorted(name for name in unresolved if name),
        "sheets": results,
    }


def ratio(numerator: int, denominator: int) -> str:
    return "n/a" if not denominator else f"{numerator}/{denominator} = {numerator / denominator:.3f}"


def markdown_summary(report: dict[str, Any]) -> str:
    lines = [
        "# Excel-to-Raw Contract Audit",
        "",
        f"Audit version: {report['audit_version']}",
        "",
        "This audit is answer-blind and does not change model outputs.",
        "",
        "## Available Sources",
        "",
        f"- Resolved workbooks: {report['resolved_workbook_count']}",
        f"- Resolved sheets: {report['resolved_sheet_count']}",
        f"- Status counts: {report['status_counts']}",
        "",
        "## Cell Retention",
        "",
        f"- Exact matches: {ratio(report['exact_match_count'], report['compared_nonempty_union_cells'])}",
        f"- Whitespace-normalized matches: {ratio(report['loose_match_count'], report['compared_nonempty_union_cells'])}",
        f"- Semantic-equivalent matches: {ratio(report['semantic_match_count'], report['compared_nonempty_union_cells'])}",
        f"- Excel non-empty but Raw empty: {report['raw_missing_count']}",
        f"- Non-whitespace value mismatches: {report['value_mismatch_count']}",
        f"- Excel non-empty cells outside Raw dimensions: {report['outside_raw_bounds_nonempty_count']}",
        "",
        "## Text Integrity",
        "",
        f"- Excel sheets with high-confidence suspicious text: {report['excel_suspicious_sheet_count']}",
        f"- Raw sheets with high-confidence suspicious text: {report['raw_suspicious_sheet_count']}",
        "",
        "## Interpretation",
        "",
        "- Terminal rendering artifacts are not counted as file corruption; checks use decoded Unicode code points.",
        "- The result applies only to source workbooks currently available. Missing workbooks remain an unresolved provenance gap.",
        "- Exact mismatch can reflect cell-type formatting differences. The loose comparison removes only whitespace and case differences; it does not invent or clinically reinterpret content.",
        "",
    ]
    mismatched = [
        item
        for item in report["sheets"]
        if item.get("raw_missing_count", 0) or item.get("value_mismatch_count", 0)
        or item.get("outside_raw_bounds_nonempty_count", 0)
    ]
    if mismatched:
        lines.extend(["## Sheets Requiring Review", ""])
        for item in mismatched:
            lines.append(
                f"- {Path(item['workbook_path']).name} / {item['sheet']}: "
                f"raw_missing={item['raw_missing_count']}, "
                f"value_mismatch={item['value_mismatch_count']}, "
                f"outside_bounds={item['outside_raw_bounds_nonempty_count']}"
            )
        lines.append("")
    return "\n".join(lines)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-root", type=Path, required=True)
    parser.add_argument("--workbook-root", action="append", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--markdown", type=Path)
    args = parser.parse_args(argv)
    report = audit(args.raw_root, args.workbook_root)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8-sig"
    )
    if args.markdown:
        args.markdown.parent.mkdir(parents=True, exist_ok=True)
        args.markdown.write_text(markdown_summary(report), encoding="utf-8-sig")
    print(
        f"workbooks={report['resolved_workbook_count']} "
        f"sheets={report['resolved_sheet_count']} "
        f"exact={report['exact_match_count']}/"
        f"{report['compared_nonempty_union_cells']} "
        f"semantic={report['semantic_match_count']}/"
        f"{report['compared_nonempty_union_cells']} "
        f"raw_missing={report['raw_missing_count']} "
        f"mismatch={report['value_mismatch_count']} "
        f"outside_bounds={report['outside_raw_bounds_nonempty_count']}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
