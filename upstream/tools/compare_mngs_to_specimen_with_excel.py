from __future__ import annotations

import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, Sequence

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from tools import pathogen_normalization as pathogen_names  # noqa: E402

DEFAULT_INPUT = Path("outputs") / "mngs_to_specimen_outputs.json"
DEFAULT_EXCEL_DIR = Path(r"C:\Users\User\Desktop\patient_info")
DEFAULT_OUTPUT = Path("outputs") / "mngs_to_specimen_vs_excel_report.json"

PATIENT_COL = 7
SPECIMEN_COL = 9
PATHOGEN_COLS = (13, 15, 16, 17)
VALUE_PATTERN = re.compile(r"^(?P<name>.+?)\s*\((?P<count>[^()]+)\)\s*$")
NAME_SANITIZER = re.compile(r"[^a-z0-9]+")
LIKELIHOOD_KEYS = ("high", "high_medium", "all_levels")
HEADER_SANITIZER = re.compile(r"[^a-z0-9\u4e00-\u9fff]+")
PATHOGEN_HEADER_PATTERN = re.compile(r"^pathogen\d+$")
REPORT_PATHOGEN_HEADER_PATTERN = re.compile(r"^reportpathogens\d+$")

COMPOUND_ALIAS_SPLITS = {
    "hhv7humangammaherpesvirus4": ["HHV-7", "Human gammaherpesvirus 4"],
    "humangammaherpesvirus4hhv7": ["Human gammaherpesvirus 4", "HHV-7"],
    "cmvebv": ["CMV", "EBV"],
    "ebvcmv": ["EBV", "CMV"],
}

PATIENT_HEADER_CANDIDATES = (
    "patient id",
    "patient_id",
    "hospital_patient_id",
    "hospitalpatientid",
    "code",
)
SPECIMEN_HEADER_CANDIDATES = (
    "specimen code",
    "specimen_code",
    "specimen id",
    "specimen_id",
    "specimenid",
    "sample code",
    "sample_code",
)
PATHOGEN_HEADER_CANDIDATES = (
    "bacterial suspected pathogen",
    "fungal suspected pathogen",
    "viral suspected pathogen",
    "parasite suspected pathogen",
)
MNGS_SHEET_HEADER_BONUS = ("mngs", "specimenid", "pathogen1")
GENUS_RELAXED_PREFIXES = (
    "serratia",
    "enterobacter",
    "elizabethkingia",
    "chryseobacterium",
    "pseudomonas",
    "staphylococcus",
    "candida",
)
FAMILY_EQUIVALENCE_PREFIXES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("humanalphaherpesvirus", ("humanalphaherpesvirus1", "humanalphaherpesvirus2")),
    ("enterobacterales", ("enterobacter", "serratia", "escherichia", "klebsiella", "citrobacter")),
    ("xanthomonadaceae", ("stenotrophomonas", "pseudoxanthomonas", "xanthomonas")),
)


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Compare mngs_to_specimen_outputs.json against the specimen Excel "
            "by specimen and patient. Supports tiered outputs (high/medium/low)."
        )
    )
    parser.add_argument("--input-json", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--specimen-excel", type=Path, default=None)
    parser.add_argument("--output-json", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args(argv)


def _resolve_default_excel() -> Path:
    xlsx_files = sorted(
        path for path in DEFAULT_EXCEL_DIR.glob("*.xlsx") if not path.name.startswith("~$")
    )
    if not xlsx_files:
        return DEFAULT_EXCEL_DIR / "specimen_lookup.xlsx"
    return xlsx_files[0]


def _normalize_name(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    normalized = pathogen_names.canonical_key(text)
    for canonical, prefixes in FAMILY_EQUIVALENCE_PREFIXES:
        if any(normalized.startswith(prefix) for prefix in prefixes):
            return canonical
    for genus_prefix in GENUS_RELAXED_PREFIXES:
        if normalized.startswith(genus_prefix):
            return genus_prefix
    return normalized


def _normalize_identifier(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if text.endswith(".0"):
        integer_part = text[:-2]
        if integer_part.isdigit():
            return integer_part
    return text


def _normalize_header_label(value: Any) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return ""
    return HEADER_SANITIZER.sub("", text)


def _resolve_column_index(
    header_map: dict[str, int],
    candidates: tuple[str, ...],
    fallback: int,
) -> int:
    for candidate in candidates:
        normalized = _normalize_header_label(candidate)
        if normalized in header_map:
            return header_map[normalized]
    return fallback


def _resolve_pathogen_columns(header_map: dict[str, int]) -> tuple[int, ...]:
    direct_pathogen_cols = sorted(
        idx for label, idx in header_map.items() if PATHOGEN_HEADER_PATTERN.match(label)
    )
    if direct_pathogen_cols:
        return tuple(direct_pathogen_cols)

    report_pathogen_cols = sorted(
        idx for label, idx in header_map.items() if REPORT_PATHOGEN_HEADER_PATTERN.match(label)
    )
    if report_pathogen_cols:
        return tuple(report_pathogen_cols)

    return tuple(
        _resolve_column_index(
            header_map,
            (candidate,),
            fallback,
        )
        for candidate, fallback in zip(PATHOGEN_HEADER_CANDIDATES, PATHOGEN_COLS)
    )


def _extract_excel_names(value: Any) -> list[str]:
    if value is None:
        return []
    text = str(value).strip()
    if not text or text == "-":
        return []
    names: list[str] = []
    expanded_text = text.replace("\r", "\n")
    for raw in re.split(r"[\n;\uFF1B\u3001]+", expanded_text):
        candidate = raw.strip()
        if not candidate or candidate == "-":
            continue
        match = VALUE_PATTERN.match(candidate)
        if match:
            candidate = match.group("name").strip()
        if not candidate:
            continue
        compound_parts = COMPOUND_ALIAS_SPLITS.get(_normalize_name(candidate))
        if compound_parts:
            names.extend(compound_parts)
            continue
        if "/" in candidate or "\uFF0F" in candidate:
            slash_parts = [
                part.strip()
                for part in candidate.replace("\uFF0F", "/").split("/")
                if part.strip()
            ]
            if len(slash_parts) > 1 and all(pathogen_names.canonical_key(part) for part in slash_parts):
                names.extend(slash_parts)
                continue
        normalized_candidate = _normalize_name(candidate)
        if normalized_candidate.isdigit():
            continue
        names.append(candidate)
    return names


def _load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig") as handle:
        payload = json.load(handle)
    if not isinstance(payload, dict):
        raise TypeError(f"Expected dict payload from {path}, got {type(payload)}")
    return payload


def _safe_row_value(row: Sequence[Any], index: int) -> Any:
    if index < 0 or index >= len(row):
        return None
    return row[index]


def _pick_excel_sheet(workbook: Any) -> Any:
    patient_candidates = {
        _normalize_header_label(candidate) for candidate in PATIENT_HEADER_CANDIDATES
    }
    specimen_candidates = {
        _normalize_header_label(candidate) for candidate in SPECIMEN_HEADER_CANDIDATES
    }
    fallback_pathogen_candidates = {
        _normalize_header_label(candidate) for candidate in PATHOGEN_HEADER_CANDIDATES
    }

    best_worksheet = None
    best_score = -1
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        header_map = {
            _normalize_header_label(value): idx
            for idx, value in enumerate(header)
            if value is not None and _normalize_header_label(value)
        }
        header_labels = set(header_map)
        sheet_label = _normalize_header_label(sheet_name)

        has_patient = bool(header_labels & patient_candidates)
        has_specimen = bool(header_labels & specimen_candidates)
        direct_pathogens = [
            idx for label, idx in header_map.items() if PATHOGEN_HEADER_PATTERN.match(label)
        ]
        report_pathogens = [
            idx for label, idx in header_map.items() if REPORT_PATHOGEN_HEADER_PATTERN.match(label)
        ]
        has_fallback_pathogens = bool(header_labels & fallback_pathogen_candidates)

        score = 0
        if "mngs" in sheet_label:
            score += 8
        if has_specimen:
            score += 12
        if has_patient:
            score += 8
        if direct_pathogens:
            score += min(12, len(direct_pathogens) * 2)
        elif report_pathogens:
            score += min(8, len(report_pathogens))
        elif has_fallback_pathogens:
            score += 6
        if all(token in header_labels for token in MNGS_SHEET_HEADER_BONUS):
            score += 12

        if score > best_score:
            best_score = score
            best_worksheet = worksheet

    if best_worksheet is not None and best_score > 0:
        return best_worksheet

    required_max_col = max(PATIENT_COL, SPECIMEN_COL, *PATHOGEN_COLS)
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if len(header) > required_max_col:
            return worksheet
    raise ValueError(
        "No worksheet has enough columns for fixed index lookup. "
        f"Need at least {required_max_col + 1} columns."
    )


def _load_excel_lookup(path: Path) -> tuple[dict[str, set[str]], dict[str, set[str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = _pick_excel_sheet(workbook)
    header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    header_map = {
        _normalize_header_label(value): idx
        for idx, value in enumerate(header)
        if value is not None and _normalize_header_label(value)
    }
    patient_col = _resolve_column_index(header_map, PATIENT_HEADER_CANDIDATES, PATIENT_COL)
    specimen_col = _resolve_column_index(header_map, SPECIMEN_HEADER_CANDIDATES, SPECIMEN_COL)
    pathogen_cols = _resolve_pathogen_columns(header_map)

    by_specimen: dict[str, set[str]] = defaultdict(set)
    by_patient: dict[str, set[str]] = defaultdict(set)
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        patient_cell = _safe_row_value(row, patient_col)
        specimen_cell = _safe_row_value(row, specimen_col)
        patient_id = _normalize_identifier(patient_cell)
        specimen_code = _normalize_identifier(specimen_cell)
        names: set[str] = set()
        for col in pathogen_cols:
            cell_value = _safe_row_value(row, col)
            names.update(
                normalized
                for normalized in (_normalize_name(name) for name in _extract_excel_names(cell_value))
                if normalized
            )
        if specimen_code:
            by_specimen[specimen_code].update(names)
        if patient_id and patient_id != "-":
            by_patient[patient_id].update(names)
    return dict(by_specimen), dict(by_patient)


def _extract_names_from_group(pathogens_group: Any) -> set[str]:
    normalized_names: set[str] = set()
    if not isinstance(pathogens_group, dict):
        return normalized_names
    for values in pathogens_group.values():
        if not isinstance(values, list):
            continue
        for item in values:
            if not isinstance(item, dict):
                continue
            normalized = _normalize_name(item.get("name"))
            if normalized:
                normalized_names.add(normalized)
    return normalized_names


def _load_output_lookup_by_level(
    path: Path,
) -> tuple[dict[str, dict[str, set[str]]], dict[str, dict[str, set[str]]]]:
    payload = _load_json(path)
    by_specimen: dict[str, dict[str, set[str]]] = {
        key: defaultdict(set) for key in LIKELIHOOD_KEYS
    }
    by_patient: dict[str, dict[str, set[str]]] = {
        key: defaultdict(set) for key in LIKELIHOOD_KEYS
    }

    for patient_id, entries in payload.items():
        if not isinstance(entries, list):
            continue
        patient_key = _normalize_identifier(patient_id)
        for entry in entries:
            if not isinstance(entry, dict):
                continue
            specimen_code = _normalize_identifier(entry.get("specimen_code", ""))
            tiered = entry.get("pathogens_by_likelihood")
            if isinstance(tiered, dict):
                high_names = _extract_names_from_group(tiered.get("high", {}))
                medium_names = _extract_names_from_group(tiered.get("medium", {}))
                low_names = _extract_names_from_group(
                    tiered.get("low_colonizer", tiered.get("low", {}))
                )
            else:
                legacy = _extract_names_from_group(entry.get("pathogens", {}))
                high_names = legacy
                medium_names = set()
                low_names = set()

            names_by_level = {
                "high": high_names,
                "high_medium": high_names | medium_names,
                "all_levels": high_names | medium_names | low_names,
            }
            for level_key, names in names_by_level.items():
                if specimen_code:
                    by_specimen[level_key][specimen_code].update(names)
                if patient_key:
                    by_patient[level_key][patient_key].update(names)

    return (
        {level: dict(values) for level, values in by_specimen.items()},
        {level: dict(values) for level, values in by_patient.items()},
    )


def _derive_excel_by_patient_from_output_specimens(
    output_json: Path,
    excel_by_specimen: dict[str, set[str]],
) -> dict[str, set[str]]:
    payload = _load_json(output_json)
    by_patient: dict[str, set[str]] = defaultdict(set)
    for patient_id, entries in payload.items():
        if not isinstance(entries, list):
            continue
        normalized_patient = _normalize_identifier(patient_id)
        if not normalized_patient:
            continue
        for entry in entries:
            if not isinstance(entry, dict):
                continue
            specimen_code = _normalize_identifier(entry.get("specimen_code", ""))
            if not specimen_code:
                continue
            by_patient[normalized_patient].update(excel_by_specimen.get(specimen_code, set()))
    return dict(by_patient)


def _safe_ratio(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 0.0
    return numerator / denominator


def _f1_score(precision: float, recall: float) -> float:
    if precision + recall == 0:
        return 0.0
    return 2 * precision * recall / (precision + recall)


def _build_single_comparison(lhs: dict[str, set[str]], rhs: dict[str, set[str]]) -> dict[str, Any]:
    shared_keys = sorted(set(lhs) & set(rhs), key=lambda value: (not str(value).isdigit(), str(value)))
    compared: dict[str, Any] = {}
    total_match = 0
    total_lhs = 0
    total_rhs = 0

    for key in shared_keys:
        left = lhs[key]
        right = rhs[key]
        matched = sorted(left & right)
        lhs_only = sorted(left - right)
        rhs_only = sorted(right - left)
        matched_count = len(matched)
        output_count = len(left)
        excel_count = len(right)

        total_match += matched_count
        total_lhs += output_count
        total_rhs += excel_count

        precision = _safe_ratio(matched_count, output_count)
        recall = _safe_ratio(matched_count, excel_count)
        compared[key] = {
            "output_count": output_count,
            "excel_count": excel_count,
            "matched_count": matched_count,
            "precision": precision,
            "recall": recall,
            "f1": _f1_score(precision, recall),
            "output_only": lhs_only,
            "excel_only": rhs_only,
            "matched": matched,
        }

    precision = _safe_ratio(total_match, total_lhs)
    recall = _safe_ratio(total_match, total_rhs)
    union = total_lhs + total_rhs - total_match
    return {
        "shared_keys": len(shared_keys),
        "output_only_keys": sorted(set(lhs) - set(rhs)),
        "excel_only_keys": sorted(set(rhs) - set(lhs)),
        "total_output_species": total_lhs,
        "total_excel_species": total_rhs,
        "total_matched_species": total_match,
        "precision": precision,
        "recall": recall,
        "f1": _f1_score(precision, recall),
        "jaccard": _safe_ratio(total_match, union),
        "details": compared,
    }


def _build_comparison_by_level(
    output_by_level: dict[str, dict[str, set[str]]],
    excel_lookup: dict[str, set[str]],
) -> dict[str, Any]:
    return {
        level_key: _build_single_comparison(output_by_level.get(level_key, {}), excel_lookup)
        for level_key in LIKELIHOOD_KEYS
    }


def _extract_scope_summary(by_level: dict[str, Any]) -> dict[str, Any]:
    summary: dict[str, Any] = {}
    for level_key in LIKELIHOOD_KEYS:
        stats = by_level.get(level_key, {})
        summary[level_key] = {
            "precision": stats.get("precision", 0.0),
            "recall": stats.get("recall", 0.0),
            "f1": stats.get("f1", 0.0),
            "jaccard": stats.get("jaccard", 0.0),
            "total_output_species": stats.get("total_output_species", 0),
            "total_excel_species": stats.get("total_excel_species", 0),
            "total_matched_species": stats.get("total_matched_species", 0),
            "shared_keys": stats.get("shared_keys", 0),
            "output_only_keys_count": len(stats.get("output_only_keys", [])),
            "excel_only_keys_count": len(stats.get("excel_only_keys", [])),
        }
    return summary


def _build_summary(by_specimen: dict[str, Any], by_patient: dict[str, Any]) -> dict[str, Any]:
    by_specimen_summary = _extract_scope_summary(by_specimen)
    by_patient_summary = _extract_scope_summary(by_patient)
    return {
        "recommended_recall_view": "all_levels",
        "by_specimen_all_levels_recall": by_specimen_summary["all_levels"]["recall"],
        "by_patient_all_levels_recall": by_patient_summary["all_levels"]["recall"],
        "by_specimen": by_specimen_summary,
        "by_patient": by_patient_summary,
    }


def main(argv: Sequence[str] | None = None) -> None:
    args = parse_args(argv)
    specimen_excel = args.specimen_excel or _resolve_default_excel()
    excel_by_specimen, excel_by_patient = _load_excel_lookup(specimen_excel)
    output_by_specimen, output_by_patient = _load_output_lookup_by_level(args.input_json)
    by_specimen_report = _build_comparison_by_level(output_by_specimen, excel_by_specimen)

    patient_alignment_mode = "excel_patient_id_direct"
    output_patient_keys = set(output_by_patient.get("all_levels", {}))
    excel_patient_keys = set(excel_by_patient)
    shared_patient_keys = len(output_patient_keys & excel_patient_keys)
    shared_specimen_keys = int(by_specimen_report.get("all_levels", {}).get("shared_keys", 0))
    if shared_patient_keys == 0 and shared_specimen_keys > 0:
        excel_by_patient = _derive_excel_by_patient_from_output_specimens(
            args.input_json,
            excel_by_specimen,
        )
        patient_alignment_mode = "derived_from_output_specimen_codes"

    by_patient_report = _build_comparison_by_level(output_by_patient, excel_by_patient)

    report = {
        "input_json": str(args.input_json),
        "specimen_excel": str(specimen_excel),
        "patient_alignment_mode": patient_alignment_mode,
        "likelihood_views": list(LIKELIHOOD_KEYS),
        "recommended_recall_view": "all_levels",
        "summary": _build_summary(by_specimen_report, by_patient_report),
        "by_specimen": by_specimen_report,
        "by_patient": by_patient_report,
    }

    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    with args.output_json.open("w", encoding="utf-8") as handle:
        json.dump(report, handle, ensure_ascii=False, indent=2)
    print(f"Wrote comparison report to {args.output_json}")


if __name__ == "__main__":
    main()
