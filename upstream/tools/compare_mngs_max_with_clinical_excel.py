from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from tools import pathogen_normalization as pathogen_names  # noqa: E402


NAME_SANITIZER = re.compile(r"[^a-z0-9]+")
SPLIT_PATTERN = re.compile(r"[\n\r;,，；、]+")

NON_PATHOGEN_TERMS = {
    "nopathogen",
    "nopatogen",
    "notinfectioncase",
    "noinfection",
    "noninfection",
    "negative",
    "none",
    "na",
}


@dataclass(frozen=True)
class NameEntry:
    raw: str
    canonical: str
    genus: str


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Compare deterministic mNGS max picked pathogens against clinical pathogen "
            "diagnosis columns in the merged three-hospital Excel file."
        )
    )
    parser.add_argument(
        "--patient-root",
        type=Path,
        default=Path("outputs") / "patient_info_rich_normalized",
    )
    parser.add_argument(
        "--ranked-mngs",
        type=Path,
        default=Path("outputs") / "mngs_candidate_microbes_rich_all_rk_ntc_opt_chosen_ranked.json",
    )
    parser.add_argument(
        "--max-output-suffix",
        default="mNGS_max_deterministic_opt_chosen_full",
    )
    parser.add_argument(
        "--clinical-excel",
        type=Path,
        default=None,
        help="Defaults to the newest root-level '*final-mNGS merged file.xlsx'.",
    )
    parser.add_argument(
        "--output-json",
        type=Path,
        default=Path("outputs") / "mngs_max_deterministic_vs_clinical_report.json",
    )
    parser.add_argument(
        "--output-csv",
        type=Path,
        default=Path("outputs") / "mngs_max_deterministic_vs_clinical_report.csv",
    )
    return parser.parse_args(argv)


def normalize_identifier(value: Any) -> str:
    text = str(value or "").strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def canonical_name(value: Any) -> str:
    text = str(value or "").strip().lower()
    if not text or text == "-":
        return ""
    normalized = pathogen_names.raw_key(text)
    if normalized in NON_PATHOGEN_TERMS:
        return ""
    return pathogen_names.canonical_key(text)


def genus_name(value: Any) -> str:
    text = str(value or "").strip()
    if not text or text == "-":
        return ""
    return pathogen_names.genus_name(text)


def extract_names(value: Any) -> list[NameEntry]:
    if value is None:
        return []
    text = str(value).strip()
    if not text or text == "-":
        return []
    entries: list[NameEntry] = []
    for part in SPLIT_PATTERN.split(text):
        raw = part.strip()
        if not raw or raw == "-":
            continue
        canonical = canonical_name(raw)
        if not canonical or canonical.isdigit():
            continue
        entries.append(NameEntry(raw=raw, canonical=canonical, genus=genus_name(raw)))
    return entries


def names_match(left: NameEntry, right: NameEntry) -> str:
    return pathogen_names.name_match_type(left.raw, right.raw)


def resolve_default_excel() -> Path:
    candidates = sorted(
        (p for p in Path(".").glob("*final-mNGS merged file.xlsx") if not p.name.startswith("~$")),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise FileNotFoundError("Cannot find root-level '*final-mNGS merged file.xlsx'.")
    return candidates[0]


def load_clinical_lookup(excel_path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[1]  # 微生物診斷資料
    header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    specimen_idx = next(
        (idx for idx, value in enumerate(header) if str(value or "").strip().lower() == "specimen_id"),
        1,
    )
    clinical_cols = [
        idx
        for idx, value in enumerate(header)
        if str(value or "").strip().startswith("臨床病原診斷")
    ]
    if not clinical_cols:
        raise ValueError("No clinical pathogen columns found: expected 臨床病原診斷1..8")

    lookup: dict[str, dict[str, Any]] = {}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        specimen_id = normalize_identifier(row[specimen_idx] if specimen_idx < len(row) else "")
        if not specimen_id:
            continue
        entries: list[NameEntry] = []
        raw_values: list[str] = []
        for col in clinical_cols:
            cell = row[col] if col < len(row) else None
            raw_values.append(str(cell).strip() if cell is not None else "")
            entries.extend(extract_names(cell))
        lookup[specimen_id] = {
            "raw_values": [value for value in raw_values if value],
            "entries": entries,
        }
    return lookup


def load_specimen_by_patient(ranked_path: Path) -> dict[str, str]:
    payload = json.loads(ranked_path.read_text(encoding="utf-8-sig"))
    patients = payload.get("patients", {})
    result: dict[str, str] = {}
    if not isinstance(patients, dict):
        return result
    for patient_id, patient_payload in patients.items():
        records = patient_payload.get("records", []) if isinstance(patient_payload, dict) else []
        if not records:
            continue
        specimen = normalize_identifier(records[0].get("specimen_code", ""))
        if specimen:
            result[normalize_identifier(patient_id)] = specimen
    return result


def load_picked_for_patient(patient_dir: Path, suffix: str) -> list[NameEntry]:
    patient_id = patient_dir.name.split("_")[2]
    path = patient_dir / "summary_outputs" / f"NGS_patient_{patient_id}_{suffix}.json"
    if not path.exists():
        return []
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    picked = payload.get("best_available_summary", {}).get("picked_pathogens", [])
    entries: list[NameEntry] = []
    for item in picked:
        if not isinstance(item, dict):
            continue
        name = item.get("organism_name")
        extracted = extract_names(name)
        if extracted:
            entries.append(extracted[0])
    return entries


def pairwise_compare(output_entries: list[NameEntry], clinical_entries: list[NameEntry]) -> dict[str, Any]:
    unmatched_clinical = set(range(len(clinical_entries)))
    matched: list[dict[str, str]] = []
    output_only: list[str] = []

    for output in output_entries:
        best_idx = None
        best_type = ""
        for idx in list(unmatched_clinical):
            match_type = names_match(output, clinical_entries[idx])
            if not match_type:
                continue
            if match_type == "exact_or_alias":
                best_idx = idx
                best_type = match_type
                break
            if best_idx is None or (
                match_type == "approved_group_member_match" and best_type == "genus_relaxed"
            ):
                best_idx = idx
                best_type = match_type
        if best_idx is None:
            output_only.append(output.raw)
            continue
        unmatched_clinical.remove(best_idx)
        matched.append(
            {
                "output": output.raw,
                "clinical": clinical_entries[best_idx].raw,
                "match_type": best_type,
            }
        )

    clinical_only = [clinical_entries[idx].raw for idx in sorted(unmatched_clinical)]
    return {
        "matched": matched,
        "output_only": output_only,
        "clinical_only": clinical_only,
    }


def safe_ratio(numerator: int, denominator: int) -> float:
    return numerator / denominator if denominator else 0.0


def f1_score(precision: float, recall: float) -> float:
    return 2 * precision * recall / (precision + recall) if precision + recall else 0.0


def main(argv: Sequence[str] | None = None) -> None:
    args = parse_args(argv)
    excel_path = args.clinical_excel or resolve_default_excel()
    clinical_by_specimen = load_clinical_lookup(excel_path)
    specimen_by_patient = load_specimen_by_patient(args.ranked_mngs)

    details: dict[str, Any] = {}
    total_output = 0
    total_clinical = 0
    total_matched = 0
    exact_or_alias = 0
    approved_group_member = 0
    genus_relaxed = 0
    missing_output: list[str] = []
    missing_clinical: list[str] = []

    patient_dirs = sorted(
        args.patient_root.glob("NGS_patient_*_json"),
        key=lambda p: int(p.name.split("_")[2]),
    )
    for patient_dir in patient_dirs:
        patient_id = normalize_identifier(patient_dir.name.split("_")[2])
        specimen_id = specimen_by_patient.get(patient_id, "")
        output_entries = load_picked_for_patient(patient_dir, args.max_output_suffix)
        if not output_entries:
            missing_output.append(patient_id)
        clinical_payload = clinical_by_specimen.get(specimen_id)
        if clinical_payload is None:
            missing_clinical.append(patient_id)
            clinical_entries: list[NameEntry] = []
            clinical_raw: list[str] = []
        else:
            clinical_entries = clinical_payload["entries"]
            clinical_raw = clinical_payload["raw_values"]

        comparison = pairwise_compare(output_entries, clinical_entries)
        matched = comparison["matched"]
        total_output += len(output_entries)
        total_clinical += len(clinical_entries)
        total_matched += len(matched)
        exact_or_alias += sum(1 for item in matched if item["match_type"] == "exact_or_alias")
        approved_group_member += sum(
            1 for item in matched if item["match_type"] == "approved_group_member_match"
        )
        genus_relaxed += sum(1 for item in matched if item["match_type"] == "genus_relaxed")

        details[patient_id] = {
            "patient_id": patient_id,
            "specimen_id": specimen_id,
            "picked_pathogens": [entry.raw for entry in output_entries],
            "clinical_pathogens": clinical_raw,
            **comparison,
            "output_count": len(output_entries),
            "clinical_count": len(clinical_entries),
            "matched_count": len(matched),
        }

    precision = safe_ratio(total_matched, total_output)
    recall = safe_ratio(total_matched, total_clinical)
    report = {
        "patient_root": str(args.patient_root),
        "ranked_mngs": str(args.ranked_mngs),
        "max_output_suffix": args.max_output_suffix,
        "clinical_excel": str(excel_path),
        "clinical_sheet": "微生物診斷資料",
        "clinical_columns": "臨床病原診斷1..8",
        "summary": {
            "patients_compared": len(details),
            "total_output_picked_pathogens": total_output,
            "total_clinical_pathogens": total_clinical,
            "total_matched_pathogens": total_matched,
            "exact_or_alias_matches": exact_or_alias,
            "approved_group_member_matches": approved_group_member,
            "genus_relaxed_matches": genus_relaxed,
            "precision": precision,
            "recall": recall,
            "f1": f1_score(precision, recall),
            "missing_output_patient_ids": missing_output,
            "missing_clinical_patient_ids": missing_clinical,
        },
        "details": details,
    }

    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    args.output_csv.parent.mkdir(parents=True, exist_ok=True)
    with args.output_csv.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "patient_id",
                "specimen_id",
                "picked_pathogens",
                "clinical_pathogens",
                "matched",
                "output_only",
                "clinical_only",
                "output_count",
                "clinical_count",
                "matched_count",
            ],
        )
        writer.writeheader()
        for row in details.values():
            writer.writerow(
                {
                    "patient_id": row["patient_id"],
                    "specimen_id": row["specimen_id"],
                    "picked_pathogens": "; ".join(row["picked_pathogens"]),
                    "clinical_pathogens": "; ".join(row["clinical_pathogens"]),
                    "matched": "; ".join(
                        f'{item["output"]} ~ {item["clinical"]} ({item["match_type"]})'
                        for item in row["matched"]
                    ),
                    "output_only": "; ".join(row["output_only"]),
                    "clinical_only": "; ".join(row["clinical_only"]),
                    "output_count": row["output_count"],
                    "clinical_count": row["clinical_count"],
                    "matched_count": row["matched_count"],
                }
            )

    print(f"Wrote JSON report: {args.output_json}")
    print(f"Wrote CSV report: {args.output_csv}")
    print(json.dumps(report["summary"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
