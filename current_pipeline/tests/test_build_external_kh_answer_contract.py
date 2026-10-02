import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from tools import build_external_kh_answer_contract as contract


class ExternalKhAnswerContractTest(unittest.TestCase):
    def test_normalize_workbook_answer_removes_context_not_identity(self):
        self.assertEqual(
            contract.normalize_workbook_answer("Enterococcus faecium (VRE)-urine"),
            ("Enterococcus faecium", False),
        )
        self.assertEqual(
            contract.normalize_workbook_answer("Klebsiella pneumoniae_plasma"),
            ("Klebsiella pneumoniae", False),
        )
        self.assertEqual(contract.normalize_workbook_answer("HSV-1"), ("HSV-1", False))
        self.assertEqual(contract.normalize_workbook_answer("No"), ("", True))

    def test_run_preserves_development_and_unblinds_only_heldout(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "answers.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.cell(1, 2, "patient_id")
            for patient_id in (1, 2, 3):
                sheet.cell(patient_id + 1, 2, patient_id)
            sheet.cell(2, 85, "Pneumocystis jirovecii")
            sheet.cell(2, 86, "Pneumocystis jirovecii")
            sheet.cell(3, 85, "No")
            sheet.cell(4, 85, "Workbook value must not replace revised answer")
            workbook.save(workbook_path)
            workbook.close()

            development_path = root / "development.csv"
            with development_path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=["patient_id", "answers", "revision_status"],
                )
                writer.writeheader()
                writer.writerow({
                    "patient_id": 3,
                    "answers": "Haemophilus influenzae",
                    "revision_status": "revised",
                })
            freeze_path = root / "freeze.json"
            freeze_path.write_text(json.dumps({
                "passed": True,
                "answers_loaded": False,
                "expected_patient_ids": [1, 2, 3],
                "cohort_partition": {"heldout_patient_ids": [1, 2]},
            }), encoding="utf-8")
            output = root / "output"
            summary = contract.run(
                workbook_path, development_path, freeze_path, output
            )
            self.assertEqual(summary["patient_count"], 3)
            self.assertEqual(summary["answer_organism_count"], 2)
            with (output / "kh_51patient_answers_postfreeze_20260930.csv").open(
                "r", encoding="utf-8-sig", newline=""
            ) as handle:
                rows = {int(row["patient_id"]): row for row in csv.DictReader(handle)}
            self.assertEqual(rows[1]["answers"], "Pneumocystis jirovecii")
            self.assertEqual(rows[2]["label_status"], "explicit_no_pathogen")
            self.assertEqual(rows[3]["answers"], "Haemophilus influenzae")

    def test_heldout_override_removes_one_answer_with_audit_note(self):
        rows = {
            42: {
                "names": ["Pseudomonas aeruginosa", "Enterococcus faecium"],
                "explicit_negative": False,
                "raw_values": ["Pseudomonas aeruginosa", "Enterococcus faecium"],
                "source_row": 42,
            }
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "overrides.json"
            path.write_text(json.dumps({
                "scope": "post-freeze benchmark correction only",
                "patients": {
                    "42": {
                        "remove": ["Enterococcus faecium"],
                        "reason": "confirmed correction",
                    }
                },
            }), encoding="utf-8")
            corrected, notes = contract.apply_heldout_overrides(rows, path, {42})
            self.assertEqual(corrected[42]["names"], ["Pseudomonas aeruginosa"])
            self.assertEqual(notes[42], "confirmed correction")


if __name__ == "__main__":
    unittest.main()
