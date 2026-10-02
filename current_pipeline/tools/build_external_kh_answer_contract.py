"""Build a post-freeze 51-patient KH benchmark answer contract.

The development cohort keeps the clinically revised 2026-09-18 answers.
Only held-out patients are read from the physician clinical-diagnosis workbook.
This tool must run after the decision freeze and writes to a separate evaluation
directory so answers cannot enter the decision pipeline.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from tools.pathogen_normalization import canonical_key


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "KMUH-mNGS study-record - 臨床病源判斷.xlsx"
DEFAULT_DEVELOPMENT_ANSWERS = (
    ROOT
    / "outputs/runs/2026-09-18_KH_answer_revision_metrics"
    / "kh_answers_clinical_revision_20260918.csv"
)
DEFAULT_FREEZE = (
    ROOT
    / "outputs/runs/2026-09-30_51patients_integrated_release_v2_6_heldout_freeze_v3"
    / "answer_blind_freeze.json"
)
DEFAULT_HELDOUT_OVERRIDES = (
    ROOT / "rules" / "kh_heldout_answer_overrides_20261001.json"
)
DEFAULT_OUTPUT = (
    ROOT
    / "outputs/runs/2026-09-30_51patients_integrated_release_v2_6_heldout_freeze_v3_evaluation"
    / "answer_contract_v2_20261001"
)

PATIENT_ID_COLUMN = 2
CLINICAL_DIAGNOSIS_COLUMNS = (85, 86, 87, 88, 89, 91, 92)
SPLIT_RE = re.compile(r"[\n\r;,，；、]+")
SPECIMEN_SUFFIX_RE = re.compile(
    r"\s*(?:\([^)]*\))?\s*[-_]\s*(?:urine|plasma|blood|serum|balf)\s*$",
    re.I,
)
VRE_SUFFIX_RE = re.compile(r"\s*\(VRE\)\s*$", re.I)
NO_PATHOGEN_VALUES = {"-", "na", "n/a", "no", "none", "negative", "no pathogen"}


def read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(value, dict):
        raise ValueError(f"Expected JSON object: {path}")
    return value


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def normalize_workbook_answer(value: Any) -> tuple[str, bool]:
    text = " ".join(str(value or "").strip().split())
    if not text:
        return "", False
    if text.casefold() in NO_PATHOGEN_VALUES:
        return "", True
    text = text.lstrip(".").strip()
    text = SPECIMEN_SUFFIX_RE.sub("", text).strip()
    text = VRE_SUFFIX_RE.sub("", text).strip()
    return text, False


def extract_cell_answers(value: Any) -> tuple[list[str], bool]:
    names: list[str] = []
    explicit_negative = False
    for part in SPLIT_RE.split(str(value or "")):
        name, negative = normalize_workbook_answer(part)
        explicit_negative = explicit_negative or negative
        if name:
            names.append(name)
    return names, explicit_negative


def unique_names(values: list[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        key = canonical_key(value)
        if key and key not in seen:
            seen.add(key)
            result.append(value)
    return result


def workbook_answers(path: Path) -> dict[int, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError(
                f"Expected one worksheet in {path}, found {len(workbook.worksheets)}"
            )
        worksheet = workbook.worksheets[0]
        result: dict[int, dict[str, Any]] = {}
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            patient_value = (
                row[PATIENT_ID_COLUMN - 1] if len(row) >= PATIENT_ID_COLUMN else None
            )
            if (
                not isinstance(patient_value, (int, float))
                or int(patient_value) != patient_value
            ):
                continue
            patient_id = int(patient_value)
            raw_values: list[str] = []
            names: list[str] = []
            explicit_negative = False
            for column in CLINICAL_DIAGNOSIS_COLUMNS:
                value = row[column - 1] if len(row) >= column else None
                if value is None or not str(value).strip():
                    continue
                raw_values.append(str(value).strip())
                extracted, negative = extract_cell_answers(value)
                names.extend(extracted)
                explicit_negative = explicit_negative or negative
            names = unique_names(names)
            if names and explicit_negative:
                raise ValueError(
                    f"P{patient_id} mixes pathogen and no-pathogen labels"
                )
            if not names and not explicit_negative:
                raise ValueError(
                    f"P{patient_id} has no clinical diagnosis label in row {row_number}"
                )
            if patient_id in result:
                raise ValueError(f"Duplicate patient P{patient_id} in workbook")
            result[patient_id] = {
                "names": names,
                "explicit_negative": explicit_negative,
                "raw_values": raw_values,
                "source_row": row_number,
            }
        return result
    finally:
        workbook.close()


def apply_heldout_overrides(
    workbook_rows: dict[int, dict[str, Any]],
    path: Path,
    heldout_ids: set[int],
) -> tuple[dict[int, dict[str, Any]], dict[int, str]]:
    payload = read_json(path)
    if payload.get("scope") != "post-freeze benchmark correction only":
        raise ValueError("Held-out overrides have the wrong scope")
    result = {
        patient_id: {
            **row,
            "names": list(row["names"]),
            "raw_values": list(row["raw_values"]),
        }
        for patient_id, row in workbook_rows.items()
    }
    notes: dict[int, str] = {}
    for patient_text, override in (payload.get("patients") or {}).items():
        patient_id = int(patient_text)
        if patient_id not in heldout_ids:
            raise ValueError(f"Override P{patient_id} is not in the held-out cohort")
        if patient_id not in result:
            raise ValueError(f"Override P{patient_id} is absent from the workbook")
        if not isinstance(override, dict):
            raise ValueError(f"Override P{patient_id} must be a JSON object")
        current = result[patient_id]
        names = list(current["names"])
        remove_names = unique_names([str(value) for value in override.get("remove") or []])
        remove_keys = {canonical_key(value) for value in remove_names}
        existing_keys = {canonical_key(value) for value in names}
        missing = remove_keys - existing_keys
        if missing:
            raise ValueError(
                f"Override P{patient_id} cannot remove absent answer keys: {sorted(missing)}"
            )
        names = [value for value in names if canonical_key(value) not in remove_keys]
        names = unique_names(names + [str(value) for value in override.get("add") or []])
        if override.get("set_no_pathogen"):
            if names:
                raise ValueError(
                    f"Override P{patient_id} cannot set no-pathogen with answer names"
                )
            current["explicit_negative"] = True
        elif names:
            current["explicit_negative"] = False
        elif not current.get("explicit_negative"):
            raise ValueError(
                f"Override P{patient_id} removed every answer without set_no_pathogen"
            )
        current["names"] = names
        notes[patient_id] = str(override.get("reason") or "held-out answer override")
    return result, notes


def run(
    workbook_path: Path,
    development_answers_path: Path,
    freeze_path: Path,
    output_dir: Path,
    heldout_overrides_path: Path | None = None,
) -> dict[str, Any]:
    if output_dir.exists() and any(output_dir.iterdir()):
        raise FileExistsError(f"Refusing to overwrite {output_dir}")
    freeze = read_json(freeze_path)
    if not freeze.get("passed") or freeze.get("answers_loaded") is not False:
        raise ValueError("Decision freeze must pass with answers_loaded=false")
    expected_ids = {int(value) for value in freeze.get("expected_patient_ids") or []}
    heldout_ids = {
        int(value)
        for value in (freeze.get("cohort_partition") or {}).get("heldout_patient_ids") or []
    }
    if not expected_ids or not heldout_ids or not heldout_ids < expected_ids:
        raise ValueError("Freeze lacks a valid expected/held-out cohort partition")
    development_ids = expected_ids - heldout_ids

    development_rows = {
        int(row["patient_id"]): row for row in read_csv(development_answers_path)
    }
    if set(development_rows) != development_ids:
        raise ValueError(
            "Development answer IDs do not match the frozen development cohort: "
            f"missing={sorted(development_ids - set(development_rows))}, "
            f"extra={sorted(set(development_rows) - development_ids)}"
        )
    workbook_rows = workbook_answers(workbook_path)
    missing_heldout = heldout_ids - set(workbook_rows)
    if missing_heldout:
        raise ValueError(f"Workbook is missing held-out answers: {sorted(missing_heldout)}")
    override_notes: dict[int, str] = {}
    if heldout_overrides_path is not None:
        workbook_rows, override_notes = apply_heldout_overrides(
            workbook_rows, heldout_overrides_path, heldout_ids
        )

    rows: list[dict[str, Any]] = []
    for patient_id in sorted(expected_ids):
        if patient_id in development_ids:
            source = development_rows[patient_id]
            answer_text = (source.get("answers") or "-").strip() or "-"
            status = "explicit_no_pathogen" if answer_text == "-" else "positive"
            rows.append({
                "patient_id": patient_id,
                "answers": answer_text,
                "raw_workbook_values": "",
                "label_status": status,
                "answer_source": "clinical_revision_20260918",
                "revision_status": source.get("revision_status") or "",
                "source_row": "",
            })
            continue
        source = workbook_rows[patient_id]
        answer_text = "; ".join(source["names"]) if source["names"] else "-"
        rows.append({
            "patient_id": patient_id,
            "answers": answer_text,
            "raw_workbook_values": " | ".join(source["raw_values"]),
            "label_status": (
                "explicit_no_pathogen" if source["explicit_negative"] else "positive"
            ),
            "answer_source": "physician_clinical_diagnosis_workbook_post_freeze",
            "revision_status": (
                "newly_unblinded_heldout_20260930"
                + (
                    f"; corrected_20261001: {override_notes[patient_id]}"
                    if patient_id in override_notes
                    else ""
                )
            ),
            "source_row": source["source_row"],
        })

    output_dir.mkdir(parents=True, exist_ok=True)
    if heldout_overrides_path is not None:
        full_name = "kh_51patient_answers_postfreeze_revised_20261001.csv"
        heldout_name = "kh_18patient_heldout_answers_postfreeze_revised_20261001.csv"
    else:
        full_name = "kh_51patient_answers_postfreeze_20260930.csv"
        heldout_name = "kh_18patient_heldout_answers_postfreeze_20260930.csv"
    full_path = output_dir / full_name
    heldout_path = output_dir / heldout_name
    write_csv(full_path, rows)
    write_csv(heldout_path, [row for row in rows if int(row["patient_id"]) in heldout_ids])
    positive_rows = [row for row in rows if row["label_status"] == "positive"]
    heldout_rows = [row for row in rows if int(row["patient_id"]) in heldout_ids]
    summary = {
        "schema_version": (
            "kh_external_answer_contract.v2"
            if heldout_overrides_path is not None
            else "kh_external_answer_contract.v1"
        ),
        "scope": "post-freeze evaluation only",
        "freeze_passed_before_unblinding": True,
        "answers_loaded_into_decision_pipeline": False,
        "patient_count": len(rows),
        "positive_patient_count": len(positive_rows),
        "explicit_no_pathogen_patient_count": len(rows) - len(positive_rows),
        "answer_organism_count": sum(
            len([value for value in str(row["answers"]).split("; ") if value != "-"])
            for row in rows
        ),
        "development_patient_ids": sorted(development_ids),
        "heldout_patient_ids": sorted(heldout_ids),
        "heldout_positive_patient_count": sum(
            row["label_status"] == "positive" for row in heldout_rows
        ),
        "heldout_explicit_no_pathogen_patient_count": sum(
            row["label_status"] == "explicit_no_pathogen" for row in heldout_rows
        ),
        "heldout_answer_organism_count": sum(
            len([value for value in str(row["answers"]).split("; ") if value != "-"])
            for row in heldout_rows
        ),
        "heldout_overrides_applied": [
            {"patient_id": patient_id, "reason": reason}
            for patient_id, reason in sorted(override_notes.items())
        ],
        "source_contracts": {
            "decision_freeze": str(freeze_path.resolve()),
            "decision_freeze_sha256": sha256_file(freeze_path),
            "development_answers": str(development_answers_path.resolve()),
            "development_answers_sha256": sha256_file(development_answers_path),
            "physician_workbook": str(workbook_path.resolve()),
            "physician_workbook_sha256": sha256_file(workbook_path),
            "worksheet_index": 0,
            "patient_id_column_1_based": PATIENT_ID_COLUMN,
            "clinical_diagnosis_columns_1_based": list(CLINICAL_DIAGNOSIS_COLUMNS),
            **(
                {
                    "heldout_answer_overrides": str(
                        heldout_overrides_path.resolve()
                    ),
                    "heldout_answer_overrides_sha256": sha256_file(
                        heldout_overrides_path
                    ),
                }
                if heldout_overrides_path is not None
                else {}
            ),
        },
        "normalization": [
            "Trim leading punctuation and surrounding whitespace.",
            "Remove trailing specimen annotations: plasma, urine, blood, serum, BALF.",
            "Remove VRE resistance annotation from the organism identity.",
            "Deduplicate aliases by the versioned pathogen normalization key.",
            "Preserve explicit No labels as labeled no-pathogen cases.",
        ],
        "outputs": {
            "full_answers": str(full_path.resolve()),
            "full_answers_sha256": sha256_file(full_path),
            "heldout_answers": str(heldout_path.resolve()),
            "heldout_answers_sha256": sha256_file(heldout_path),
        },
    }
    manifest_path = output_dir / "answer_contract_manifest.json"
    manifest_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument(
        "--development-answers", type=Path, default=DEFAULT_DEVELOPMENT_ANSWERS
    )
    parser.add_argument("--freeze", type=Path, default=DEFAULT_FREEZE)
    parser.add_argument(
        "--heldout-overrides", type=Path, default=DEFAULT_HELDOUT_OVERRIDES
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    print(json.dumps(run(
        args.workbook,
        args.development_answers,
        args.freeze,
        args.output_dir,
        args.heldout_overrides,
    ), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
