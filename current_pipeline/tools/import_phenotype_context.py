"""Import physician-provided pneumonia phenotypes into answer-blind context JSON.

The importer deliberately separates three concerns:

* compact model context for downstream reasoning;
* complete normalized evidence for provenance and audit;
* interpretation gates that prevent a phenotype from becoming a pathogen label.

It does not modify deterministic tiers, LLM review, merge, or benchmark answers.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


SCHEMA_VERSION = "phenotype_context_v1.0"
FULL_EVIDENCE_SCHEMA_VERSION = "phenotype_source_evidence_v1.0"

PHENOTYPE_GROUPS = {
    "aspiration_and_host": (
        "ASPIRATION_RISK",
        "ANAEROBIC_RISK",
        "IMMUNOCOMPROMISED",
        "FUNGAL_INFECTION_RISK",
    ),
    "healthcare_and_resistance_risk": (
        "RECENT_HOSPITALIZATION_90D",
        "RECENT_IV_ANTIBIOTICS_90D",
        "PRIOR_MDR_COLONIZATION_OR_INFECTION",
        "MDR_GNB_RISK",
        "PSEUDOMONAS_RISK",
        "MRSA_RISK",
    ),
    "lung_airway_and_setting": (
        "STRUCTURAL_LUNG_DISEASE",
        "TRACHEOSTOMY_OR_CHRONIC_AIRWAY",
        "HAP_PHENOTYPE",
        "VAP_PHENOTYPE",
        "POST_OBSTRUCTIVE_PNEUMONIA",
    ),
    "severity_and_complications": (
        "SEPTIC_SHOCK",
        "RESPIRATORY_FAILURE",
        "MECHANICAL_VENTILATION",
        "ARDS",
        "RENAL_REPLACEMENT_THERAPY",
        "COMPLICATED_PNEUMONIA",
        "BACTEREMIA",
    ),
}
EXPECTED_PHENOTYPES = tuple(
    phenotype for group in PHENOTYPE_GROUPS.values() for phenotype in group
)
REQUIRED_SHEETS = (
    "Phenotypes_Wide",
    "Phenotypes_Long",
    "Review_Only",
    "Microbiology",
    "Clinical_Timeline",
    "Lab_Results",
    "Evidence",
    "Clinical_Facts",
)

RESPONSE_PATTERNS = (
    re.compile(r"\b(?:clinical(?:ly)?|symptoms?|condition)\s+(?:was|were|is|are|has|had)?\s*improv", re.I),
    re.compile(r"\bimprov(?:ed|ement|ing)\s+(?:after|with|following)\b", re.I),
    re.compile(r"\brespond(?:ed|ing|s)?\s+to\b", re.I),
    re.compile(r"\b(?:good|favorable|poor|no)\s+(?:clinical\s+)?response\s+to\b", re.I),
    re.compile(r"\b(?:resolved|resolution)\s+(?:after|with|following)\b", re.I),
    re.compile(r"\bdefervesc", re.I),
    re.compile(r"\brespiratory failure resolved\b", re.I),
)
EXCLUDED_RESPONSE_PATTERNS = (
    re.compile(r"\bno (?:light reflex nor other )?response to stimulation\b", re.I),
)
ANTIMICROBIAL_PATTERN = re.compile(
    r"\b(?:antibiotic|antimicrobial|antifungal|antiviral|meropenem|imipenem|ertapenem|"
    r"cef(?:epime|triaxone|tazidime|azolin|operazone)|piperacillin|tazobactam|vancomycin|"
    r"linezolid|levofloxacin|ciprofloxacin|azithromycin|doxycycline|trimethoprim|"
    r"sulfamethoxazole|bactrim|sevatrim|voriconazole|fluconazole|micafungin|"
    r"amphotericin|ganciclovir|valganciclovir|acyclovir|remdesivir)\b",
    re.I,
)
RELEVANT_LAB_PATTERN = re.compile(
    r"(?:^|\b)(?:WBC|white blood|CRP|C-reactive|procalcitonin|PCT|lactate|"
    r"lymph|absolute lymph|neutro|ANC|platelet|PLT)(?:\b|$)",
    re.I,
)


def json_scalar(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value.strip() if isinstance(value, str) else value
    return str(value)


def sheet_records(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration:
        return []
    headers = [str(value or "").strip() for value in header_values]
    records: list[dict[str, Any]] = []
    for source_row, values in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        record = {
            header: json_scalar(value)
            for header, value in zip(headers, values)
            if header
        }
        record["_source_row"] = source_row
        records.append(record)
    return records


def patient_number(value: Any) -> int | None:
    match = re.search(r"(?i)patient\D*(\d+)", str(value or ""))
    return int(match.group(1)) if match else None


def patient_slug(patient_id: str) -> str:
    number = patient_number(patient_id)
    return f"patient_{number}" if number is not None else re.sub(r"\W+", "_", patient_id.lower()).strip("_")


def split_reason_codes(value: Any) -> list[str]:
    return [part.strip() for part in re.split(r"[|;]", str(value or "")) if part.strip()]


def phenotype_model_role(row: dict[str, Any]) -> str:
    status = str(row.get("status") or "").upper()
    confidence = str(row.get("confidence") or "").upper()
    temporal = str(row.get("temporal_relation") or "").upper()
    qc = str(row.get("phenotype_qc_status") or "").upper()
    if status != "YES":
        return "not_established"
    if qc != "PASS" or confidence not in {"HIGH", "MEDIUM"}:
        return "manual_review"
    if temporal in {"PRE_PNEUMONIA", "AT_ONSET"}:
        return "etiologic_prior"
    if temporal == "AFTER_ONSET":
        return "outcome_or_context"
    return "context_only"


def phenotype_records(rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for raw in rows:
        phenotype = str(raw.get("phenotype") or "").strip()
        if not phenotype:
            continue
        role = phenotype_model_role(raw)
        result.append(
            {
                "phenotype": phenotype,
                "group": next(
                    (name for name, members in PHENOTYPE_GROUPS.items() if phenotype in members),
                    "unrecognized",
                ),
                "status": str(raw.get("status") or "").upper() or None,
                "confidence": str(raw.get("confidence") or "").upper() or None,
                "temporal_relation": str(raw.get("temporal_relation") or "").upper() or None,
                "phenotype_qc_status": str(raw.get("phenotype_qc_status") or "").upper() or None,
                "model_role": role,
                "can_modify_pathogen_prior": role == "etiologic_prior",
                "reason_codes": split_reason_codes(raw.get("reason_codes")),
                "reason": raw.get("reason"),
                "evidence": raw.get("evidence"),
                "evidence_page": raw.get("evidence_page"),
                "source_type": raw.get("source_type"),
                "source_date": raw.get("source_date"),
                "review_reason": raw.get("review_reason"),
                "source_row": raw.get("_source_row"),
            }
        )
    return result


def text_of(record: dict[str, Any]) -> str:
    return " ".join(str(value or "") for key, value in record.items() if key != "_source_row")


def is_explicit_response(text: str) -> bool:
    if any(pattern.search(text) for pattern in EXCLUDED_RESPONSE_PATTERNS):
        return False
    return any(pattern.search(text) for pattern in RESPONSE_PATTERNS)


def relevant_labs(rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    selected: list[dict[str, Any]] = []
    for row in rows:
        label = f"{row.get('test_name') or ''} {row.get('raw_label') or ''}"
        if RELEVANT_LAB_PATTERN.search(label):
            selected.append(row)
    return selected


def treatment_context(
    clinical_facts: list[dict[str, Any]],
    clinical_timeline: list[dict[str, Any]],
    labs: list[dict[str, Any]],
) -> dict[str, Any]:
    source_rows = clinical_facts + clinical_timeline
    treatment_events = [
        row
        for row in source_rows
        if str(row.get("fact_type") or row.get("event_type") or "").upper()
        in {"TREATMENT", "MEDICATION", "ANTIMICROBIAL"}
    ]
    response_mentions = [row for row in source_rows if is_explicit_response(text_of(row))]
    linked_antimicrobial = [
        row for row in response_mentions if ANTIMICROBIAL_PATTERN.search(text_of(row))
    ]
    selected_labs = relevant_labs(labs)
    return {
        "treatment_events": treatment_events,
        "explicit_response_mentions": response_mentions,
        "explicit_antimicrobial_response_mentions": linked_antimicrobial,
        "explicit_antimicrobial_response_available": bool(linked_antimicrobial),
        "inflammatory_and_immune_lab_series": selected_labs,
        "interpretation_constraints": [
            "CBC, CRP, procalcitonin, and lactate trends do not by themselves prove response to a specific antimicrobial.",
            "A treatment-response claim requires an explicit clinical statement or a time-linked analysis that identifies the therapy, target interval, and competing interventions.",
            "Treatment started after pneumonia onset is outcome context, not baseline pathogen evidence.",
        ],
    }


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def csv_rows(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def rows_without_source_index(rows: Iterable[dict[str, Any]]) -> list[dict[str, str]]:
    return [
        {key: "" if value is None else str(value) for key, value in row.items() if key != "_source_row"}
        for row in rows
    ]


def csv_crosscheck(
    patient_id: str,
    sheets: dict[str, list[dict[str, Any]]],
    csv_maps: dict[str, dict[int, Path]],
) -> dict[str, Any]:
    number = patient_number(patient_id)
    checks: dict[str, Any] = {}
    for sheet_name, kind in (("Clinical_Facts", "clinical_facts"), ("Microbiology", "microbiology")):
        csv_path = csv_maps.get(kind, {}).get(number or -1)
        if csv_path is None:
            checks[kind] = {"status": "missing_csv"}
            continue
        workbook_rows = rows_without_source_index(sheets[sheet_name])
        external_rows = rows_without_source_index(csv_rows(csv_path))
        checks[kind] = {
            "status": "match" if workbook_rows == external_rows else "mismatch",
            "workbook_rows": len(workbook_rows),
            "csv_rows": len(external_rows),
            "csv_file": csv_path.name,
        }
    return checks


def build_csv_maps(csv_root: Path | None) -> dict[str, dict[int, Path]]:
    maps: dict[str, dict[int, Path]] = {"clinical_facts": {}, "microbiology": {}}
    if csv_root is None or not csv_root.exists():
        return maps
    for kind in maps:
        for path in csv_root.glob(f"*_{kind}.csv"):
            number = patient_number(path.name)
            if number is not None:
                maps[kind][number] = path
    return maps


def import_workbook(path: Path, output_root: Path, csv_maps: dict[str, dict[int, Path]]) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    missing_sheets = [sheet for sheet in REQUIRED_SHEETS if sheet not in workbook.sheetnames]
    if missing_sheets:
        raise ValueError(f"{path.name}: missing sheets {missing_sheets}")
    sheets = {sheet: sheet_records(workbook, sheet) for sheet in REQUIRED_SHEETS}
    workbook.close()

    long_rows = sheets["Phenotypes_Long"]
    patient_id = str(
        (long_rows[0].get("patient_id") if long_rows else None)
        or (sheets["Phenotypes_Wide"][0].get("patient_id") if sheets["Phenotypes_Wide"] else None)
        or path.stem
    )
    slug = patient_slug(patient_id)
    phenotypes = phenotype_records(long_rows)
    observed = {item["phenotype"] for item in phenotypes}
    missing_phenotypes = sorted(set(EXPECTED_PHENOTYPES) - observed)
    extra_phenotypes = sorted(observed - set(EXPECTED_PHENOTYPES))
    wide = sheets["Phenotypes_Wide"][0] if sheets["Phenotypes_Wide"] else {}
    source_full_name = f"{slug}_phenotype_evidence_full.json"

    treatment = treatment_context(
        sheets["Clinical_Facts"], sheets["Clinical_Timeline"], sheets["Lab_Results"]
    )
    roles = Counter(item["model_role"] for item in phenotypes)
    context = {
        "schema_version": SCHEMA_VERSION,
        "patient_id": patient_id,
        "patient_number": patient_number(patient_id),
        "source": {
            "workbook_file": path.name,
            "workbook_sha256": file_sha256(path),
            "full_evidence_json": source_full_name,
            "patient_qc_status": wide.get("patient_qc_status"),
            "system_issues": str(wide.get("system_issues") or "").split(" | ") if wide.get("system_issues") else [],
            "csv_crosscheck": csv_crosscheck(patient_id, sheets, csv_maps),
        },
        "phenotype_contract": {
            "expected_count": len(EXPECTED_PHENOTYPES),
            "observed_count": len(phenotypes),
            "missing_phenotypes": missing_phenotypes,
            "unrecognized_phenotypes": extra_phenotypes,
            "model_role_counts": dict(sorted(roles.items())),
        },
        "phenotypes": phenotypes,
        "model_context": {
            "etiologic_priors": [item for item in phenotypes if item["model_role"] == "etiologic_prior"],
            "context_only": [
                item
                for item in phenotypes
                if item["model_role"] in {"context_only", "outcome_or_context"}
            ],
            "manual_review": [item for item in phenotypes if item["model_role"] == "manual_review"],
            "not_established": [item["phenotype"] for item in phenotypes if item["model_role"] == "not_established"],
            "constraints": [
                "Phenotypes modify patient-level prior plausibility and never directly select an organism.",
                "NO means not established in the supplied record and is not strong negative evidence.",
                "AFTER_ONSET evidence is outcome/context and must not be used as a baseline etiologic prior.",
                "REVIEW_REQUIRED or low-confidence evidence requires human or adjudicator review.",
            ],
        },
        "clinical_context": {
            "microbiology": sheets["Microbiology"],
            "clinical_timeline": sheets["Clinical_Timeline"],
            "clinical_facts": sheets["Clinical_Facts"],
            "treatment_and_response": treatment,
        },
    }
    full_evidence = {
        "schema_version": FULL_EVIDENCE_SCHEMA_VERSION,
        "patient_id": patient_id,
        "source_workbook": path.name,
        "source_workbook_sha256": context["source"]["workbook_sha256"],
        "sheets": sheets,
    }
    output_root.mkdir(parents=True, exist_ok=True)
    (output_root / f"{slug}_phenotype_context.json").write_text(
        json.dumps(context, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (output_root / source_full_name).write_text(
        json.dumps(full_evidence, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return {
        "patient_id": patient_id,
        "patient_number": patient_number(patient_id),
        "source_workbook": path.name,
        "context_file": f"{slug}_phenotype_context.json",
        "full_evidence_file": source_full_name,
        "phenotype_count": len(phenotypes),
        "missing_phenotypes": missing_phenotypes,
        "unrecognized_phenotypes": extra_phenotypes,
        "model_role_counts": dict(sorted(roles.items())),
        "patient_qc_status": wide.get("patient_qc_status"),
        "phenotype_review_required_count": sum(
            1 for item in phenotypes if item["phenotype_qc_status"] != "PASS"
        ),
        "treatment_event_count": len(treatment["treatment_events"]),
        "explicit_response_mention_count": len(treatment["explicit_response_mentions"]),
        "explicit_antimicrobial_response_mention_count": len(
            treatment["explicit_antimicrobial_response_mentions"]
        ),
        "csv_crosscheck": context["source"]["csv_crosscheck"],
    }


def write_readme(output_root: Path, records: list[dict[str, Any]]) -> None:
    text = f"""# Phenotype context import

This folder contains answer-blind context imported from the physician-provided MedicalPhenotype workbooks.

- Patients imported: {len(records)}
- Compact model files: `patient_<N>_phenotype_context.json`
- Full normalized provenance: `patient_<N>_phenotype_evidence_full.json`
- Cohort audit: `import_audit.json`

The compact files may modify prior plausibility only. They do not select organisms, contain benchmark answers, or alter the current scorer/merge baseline. A phenotype `NO` means that the source record did not establish the phenotype; it is not treated as strong negative evidence.

Treatment events, explicit response statements, and inflammatory/CBC lab series are kept separately. Lab trends are not interpreted as response to a specific antimicrobial without an explicit time-linked clinical statement.
"""
    (output_root / "README.md").write_text(text, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_root", type=Path, help="Directory containing *_phenotype.xlsx files")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--csv-root", type=Path, help="Optional directory containing duplicated Clinical_Facts and Microbiology CSVs")
    args = parser.parse_args()

    workbooks = sorted(
        args.input_root.glob("*_phenotype.xlsx"),
        key=lambda path: (patient_number(path.name) is None, patient_number(path.name) or 0, path.name),
    )
    if not workbooks:
        raise SystemExit(f"No *_phenotype.xlsx files found under {args.input_root}")
    csv_maps = build_csv_maps(args.csv_root)
    records = [import_workbook(path, args.output, csv_maps) for path in workbooks]
    numbers = sorted(item["patient_number"] for item in records if item["patient_number"] is not None)
    missing_numbers = (
        sorted(set(range(min(numbers), max(numbers) + 1)) - set(numbers)) if numbers else []
    )
    audit = {
        "schema_version": SCHEMA_VERSION,
        "input_root": str(args.input_root.resolve()),
        "csv_root": str(args.csv_root.resolve()) if args.csv_root else None,
        "output_root": str(args.output.resolve()),
        "patient_count": len(records),
        "patient_numbers": numbers,
        "missing_numbers_within_observed_range": missing_numbers,
        "patients_with_all_22_phenotypes": sum(
            1 for item in records if not item["missing_phenotypes"] and not item["unrecognized_phenotypes"]
        ),
        "patients_with_explicit_response_mentions": sum(
            1 for item in records if item["explicit_response_mention_count"]
        ),
        "patients_with_explicit_antimicrobial_response_mentions": sum(
            1 for item in records if item["explicit_antimicrobial_response_mention_count"]
        ),
        "csv_crosscheck_status_counts": dict(
            Counter(
                check["status"]
                for item in records
                for check in item["csv_crosscheck"].values()
            )
        ),
        "records": records,
    }
    args.output.mkdir(parents=True, exist_ok=True)
    (args.output / "import_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    write_readme(args.output, records)
    print(json.dumps({key: value for key, value in audit.items() if key != "records"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
